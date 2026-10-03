"""Muuntaa VILPE Vantaan hackathon-aineiston prototyypin datatiedostoksi.

Käyttö:
    pip install openpyxl pymupdf pillow
    python scripts/build_data.py ../Vilpe_Materiaalit

Tuottaa:
    app/data.js            window.VANTAA = {...}
    app/assets/roof.jpg    kattokartta (rajattu site_layout-PDF:stä)
"""
import glob
import json
import os
import re
import statistics as st
import sys
from collections import defaultdict
from datetime import datetime

import fitz
import openpyxl
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "..", "Vilpe_Materiaalit")
OUT = os.path.join(ROOT, "app")

# Kattokartan rajaus PDF-pisteinä (sivu 0), kattaa anturialueen.
CROP = fitz.Rect(585, 195, 1640, 700)
DPI = 130


def r1(x, n=1):
    return None if x is None else round(x, n)


def mean(xs):
    xs = [x for x in xs if x is not None]
    return st.mean(xs) if xs else None


# ---------- Kattokartta ja anturien sijainnit ----------
pdf = fitz.open(glob.glob(os.path.join(SRC, "site_layout_*.pdf"))[0])
page = pdf[0]
positions = {}
for x0, y0, x1, y1, word, *_ in page.get_text("words"):
    if re.fullmatch(r"[NP]\w{10}", word):
        # Merkki-ikoni on tekstin vasemmalla puolella.
        cx, cy = x0 - 5, (y0 + y1) / 2
        positions[word] = (
            round((cx - CROP.x0) / CROP.width, 4),
            round((cy - CROP.y0) / CROP.height, 4),
        )
pix = page.get_pixmap(dpi=DPI, clip=CROP)
Image.frombytes("RGB", (pix.width, pix.height), pix.samples).save(
    os.path.join(OUT, "assets", "roof.jpg"), quality=82)

# ---------- Vuotoanturit (RHT-2) ----------
ws = openpyxl.load_workbook(glob.glob(os.path.join(SRC, "measurements_*.xlsx"))[0], read_only=True).worksheets[0]
rows = ws.iter_rows(values_only=True)
next(rows)
daily = defaultdict(lambda: defaultdict(list))  # serial -> day -> [(t, rh, ah)]
monthly_all = defaultdict(list)
last = {}
for ts, _sid, serial, _n, _b, _bs, _bn, t, rh, ah in rows:
    if ts is None or rh is None:
        continue
    day = ts.strftime("%Y-%m-%d")
    daily[serial][day].append((t, rh, ah))
    monthly_all[ts.strftime("%Y-%m")].append((t, rh, ah))
    if serial not in last or ts > last[serial][0]:
        last[serial] = (ts, t, rh, ah)

days = sorted({d for s in daily.values() for d in s})
sensors = []
for serial in sorted(daily):
    d = daily[serial]
    allv = [v for vs in d.values() for v in vs]
    sensors.append({
        "id": serial,
        "pos": positions.get(serial),
        "rh": [r1(mean([v[1] for v in d[day]])) if day in d else None for day in days],
        "t": [r1(mean([v[0] for v in d[day]])) if day in d else None for day in days],
        "rhMean": r1(mean([v[1] for v in allv])),
        "rhMax": r1(max(v[1] for v in allv)),
        "tMin": r1(min(v[0] for v in allv if v[0] is not None)),
        "n": len(allv),
        "last": {"ts": last[serial][0].isoformat(), "t": last[serial][1], "rh": last[serial][2], "ah": last[serial][3]},
    })
network_monthly = [
    {"m": m, "t": r1(mean([v[0] for v in vs])), "rh": r1(mean([v[1] for v in vs])), "ah": r1(mean([v[2] for v in vs]))}
    for m, vs in sorted(monthly_all.items())
]

# ---------- Huippuimurit (MCU-2) ----------
COLS = {
    "Rpm": "rpm",
    "Ohjaava sisälähetin T (°C)": "tIn", "Ohjaava sisälähetin H (%)": "rhIn", "Ohjaava sisälähetin g/m³": "ahIn",
    "Ohjaava ulkolähetin T (°C)": "tOut", "Ohjaava ulkolähetin H (%)": "rhOut", "Ohjaava ulkolähetin g/m³": "ahOut",
}
units = []
for f in sorted(glob.glob(os.path.join(SRC, "VILPE Vantaa, *.xlsx"))):
    rows = list(openpyxl.load_workbook(f, read_only=True).worksheets[0].iter_rows(values_only=True))
    meta = {r[0]: r[1] for r in rows[:9] if r}
    hi = next(i for i, r in enumerate(rows) if r and r[0] == "Aikaleima" and "Rpm" in r)
    header = [COLS.get(h) for h in rows[hi]]
    recs = []
    for r in rows[hi + 1:]:
        if not r or not r[0]:
            continue
        rec = {k: v for k, v in zip(header, r) if k}
        rec["ts"] = datetime.strptime(r[0], "%Y/%m/%d %H:%M:%S")
        recs.append(rec)
    by_day, by_month = defaultdict(list), defaultdict(list)
    for rec in recs:
        by_day[rec["ts"].strftime("%Y-%m-%d")].append(rec)
        by_month[rec["ts"].strftime("%Y-%m")].append(rec)

    def agg(rs):
        return {
            "rpm": r1(mean([x["rpm"] for x in rs]), 0),
            "on": r1(100 * sum(1 for x in rs if (x["rpm"] or 0) > 0) / len(rs), 0),
            **{k: r1(mean([x.get(k) for x in rs])) for k in ("tIn", "rhIn", "ahIn", "tOut", "rhOut", "ahOut")},
        }

    # Pisin yhtäjaksoinen seisokki (rpm = 0)
    longest, cur_start, best = 0, None, (None, None)
    for rec in recs:
        if (rec["rpm"] or 0) == 0:
            cur_start = cur_start or rec["ts"]
            span = (rec["ts"] - cur_start).total_seconds() / 86400
            if span > longest:
                longest, best = span, (cur_start, rec["ts"])
        else:
            cur_start = None
    rh_in = [x.get("rhIn") for x in recs if x.get("rhIn") is not None]
    name = meta["Tunniste"].replace("VILPE Vantaa, ", "")
    units.append({
        "name": name,
        "serial": meta["Sarjanumero"],
        "purpose": meta.get("Käyttötarkoitus"),
        "mold": meta.get("Viimeisin homeindeksi"),
        "rpmLast": meta.get("Viimeisin rpm"),
        "pos": positions.get(meta["Sarjanumero"]),
        "on": r1(100 * sum(1 for x in recs if (x["rpm"] or 0) > 0) / len(recs), 0),
        "rpmMean": r1(mean([x["rpm"] for x in recs if (x["rpm"] or 0) > 0]), 0),
        "rhInMean": r1(mean(rh_in)),
        "rhIn90": r1(100 * sum(1 for x in rh_in if x > 90) / len(rh_in), 0),
        "stopDays": r1(longest, 0),
        "stopFrom": best[0].strftime("%Y-%m-%d") if best[0] else None,
        "stopTo": best[1].strftime("%Y-%m-%d") if best[1] else None,
        "days": [{"d": k, **agg(v)} for k, v in sorted(by_day.items())],
        "months": [{"m": k, **agg(v)} for k, v in sorted(by_month.items())],
    })

data = {
    "site": {"name": "VILPE Express Store, Vantaa", "structure": "Betoni · tasakatto + viherkatto + ryömintätila"},
    "days": days,
    "sensors": sensors,
    "networkMonthly": network_monthly,
    "units": units,
    "roof": {"w": pix.width, "h": pix.height},
}
with open(os.path.join(OUT, "data.js"), "w", encoding="utf-8") as fh:
    fh.write("// Generoitu: scripts/build_data.py – älä muokkaa käsin.\nwindow.VANTAA = ")
    json.dump(data, fh, ensure_ascii=False, separators=(",", ":"))
    fh.write(";\n")

print(f"{len(sensors)} anturia, {len(days)} päivää, {len(units)} imuria, "
      f"sijainteja {sum(1 for s in sensors if s['pos'])}/{len(sensors)}")
for u in units:
    print(f"  {u['name']:16} on {u['on']:>3}%  mold {u['mold']}  stop {u['stopDays']} d {u['stopFrom']}–{u['stopTo']}")
